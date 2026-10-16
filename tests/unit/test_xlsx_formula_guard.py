from openpyxl import Workbook

from pqcscan.renderers._xlsx_safe import force_text


def test_formula_strings_are_written_as_text():
    ws = Workbook().active
    ws.append(["=cmd|' /C calc'!A0", "=SUM(1)", "plain", 3])
    force_text(ws[1])
    assert [c.data_type for c in ws[1]] == ["s", "s", "s", "n"]
    assert ws["A1"].value == "=cmd|' /C calc'!A0"
