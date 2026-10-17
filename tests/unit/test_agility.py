from types import SimpleNamespace

import pytest

from pqcscan.core.agility import finding_agility, score_findings


def test_rule_table():
    assert finding_agility("fs.conf.nginx", {}) == 90
    assert finding_agility("net.tls.https", {}) == 80
    assert finding_agility("fs.cert.x509", {}) == 70
    assert finding_agility("sbom.lang.pip", {}) == 50
    assert finding_agility("code.ts.java", {}) == 30
    assert finding_agility("fs.binary.crypto", {"detection": "constant-signature"}) == 20
    assert finding_agility("fs.binary.crypto", {"linked": "linked"}) == 50
    assert finding_agility("hw.tpm_algorithms", {}) == 10
    assert finding_agility("ot.modbus.tcp", {}) == 10
    assert finding_agility("something.new", {}) == 50


def test_score_findings_weights_vulnerable_only():
    rows = [SimpleNamespace(probe_id="fs.conf.nginx", evidence={}, classification="sangat-tinggi"),
            SimpleNamespace(probe_id="ot.modbus.tcp", evidence={}, classification="sangat-tinggi"),
            SimpleNamespace(probe_id="host.platform_info", evidence={}, classification="info")]
    r = score_findings(rows)
    assert r.score == 50 and r.band == "medium" and r.by_source == {"config": 1, "hardware": 1}
    assert score_findings([]).score == 100 and score_findings([]).band == "high"


def test_fleet_drift_and_fleet_attribution_ignored():
    rows = [SimpleNamespace(probe_id="fs.conf.nginx", evidence={"_fleet": {"host": "a"}}, classification="tinggi"),
            SimpleNamespace(probe_id="fleet.drift", evidence={}, classification="tinggi")]
    r = score_findings(rows)
    assert r.score == 90 and r.by_source == {"config": 1}


def test_band_thresholds_and_ibmi_sysval():
    assert finding_agility("fs.conf.ibmi_sysval", {}) == 10
    for probe, band in (("fs.cert.x509", "high"), ("sbom.lang.pip", "medium"), ("code.ts.java", "low")):
        r = score_findings([SimpleNamespace(probe_id=probe, evidence={}, classification="tinggi")])
        assert r.band == band


def _seeded(repo):
    from pqcscan.core.types import Classification, Finding, Severity
    sid = repo.create_scan(mode="user", probe_versions={}, tool_versions={})
    repo.record_finding(sid, Finding(
        probe_id="net.tls.https", algorithm="RSA-2048", classification=Classification.SANGAT_TINGGI,
        severity=Severity.CRIT, title="rsa", evidence={"endpoint": "h:443"}))
    repo.finish_scan(sid, status="done")
    return sid


def test_report_context_carries_agility(tmp_db_path):
    from pqcscan.renderers._report_context import build_report_context
    from pqcscan.store.repo import Repo
    repo = Repo(tmp_db_path)
    repo.init_schema()
    ctx = build_report_context(repo, _seeded(repo))
    assert ctx["agility"] == {"score": 80, "band": "high", "by_source": {"server-config": 1}}


def test_bukukerja_cbom_column_is_numeric_agility(tmp_db_path, tmp_path):
    from openpyxl import load_workbook

    from pqcscan.renderers.xlsx_bukukerja import render_xlsx_bukukerja
    from pqcscan.store.repo import Repo
    repo = Repo(tmp_db_path)
    repo.init_schema()
    sid = _seeded(repo)
    for loc in ("ms", "en"):
        out = tmp_path / f"{loc}.xlsx"
        render_xlsx_bukukerja(repo, sid, out, locale=loc)
        rows = list(load_workbook(out)["2_CBOM"].iter_rows(min_row=5, values_only=True))
        assert [r[7] for r in rows if r[0]] == ["80/100"]


def test_pdf_html_shows_agility_line(tmp_db_path):
    from pqcscan.renderers.pdf_executive import build_html_executive
    from pqcscan.renderers.pdf_technical import build_html_technical
    from pqcscan.store.repo import Repo
    repo = Repo(tmp_db_path)
    repo.init_schema()
    sid = _seeded(repo)
    for build, lang, label in ((build_html_executive, "en", "Crypto-agility: 80/100 (high)"),
                               (build_html_executive, "ms", "Ketangkasan kripto: 80/100 (tinggi)"),
                               (build_html_technical, "en", "Crypto-agility: 80/100 (high)")):
        assert label in build(repo, sid, lang)


_PROBE = {10: "hw.x", 20: "storage.x", 30: "code.x", 50: "sbom.x", 70: "fs.cert.x", 80: "net.x", 90: "fs.conf.x"}


def _mix(*points):
    return [SimpleNamespace(probe_id=_PROBE[p], evidence={}, classification="tinggi") for p in points]


def test_band_edges():
    cases = [
        ((10, 10, 10, 10, 50, 90, 90), 39, "low"),      # mean 38.57
        ((10, 70), 40, "medium"),
        ((10, 20, 90, 90, 90, 90, 90), 69, "medium"),   # mean 68.57
        ((50, 90), 70, "high"),
    ]
    for points, score, band in cases:
        r = score_findings(_mix(*points))
        assert (r.score, r.band) == (score, band), points


@pytest.mark.parametrize("probe,evidence,points", [
    ("fs.conf.ibmi_sysval", {}, 10),   # must beat the fs.conf. prefix
    ("hw.tpm_algorithms", {}, 10), ("host.tpm.pcr", {}, 10), ("ot.modbus.tcp", {}, 10),
    ("storage.disk", {}, 20), ("vpn.ipsec", {}, 20),
    ("fs.binary.crypto", {"detection": "constant-signature"}, 20),
    ("fs.binary.crypto", {"linked": "embedded"}, 20),
    ("fs.binary.crypto", {"linked": "linked"}, 50),
    ("fs.binary.crypto", {}, 50),
    ("code.py.rsa", {}, 30),
    ("sbom.lang.pip", {}, 50), ("container.image", {}, 50),
    ("fs.cert.x509", {}, 70), ("fs.keystore.jks", {}, 70), ("email.smime", {}, 70), ("dns.dnssec", {}, 70),
    ("sign.codesign", {}, 70), ("trust.store", {}, 70), ("host.cloud_kms", {}, 70), ("fs.keyref.cloud", {}, 70),
    ("net.tls.https", {}, 80),
    ("host.crypto_policies", {}, 90), ("host.java.security", {}, 90), ("host.openssl.cnf", {}, 90),
    ("fs.conf.nginx", {}, 90), ("app.django", {}, 90), ("host.ssh.sshd", {}, 90), ("host.gnutls", {}, 90),
    ("host.nss", {}, 90), ("host.windows.schannel", {}, 90), ("k8s.ingress", {}, 90), ("fs.iac.terraform", {}, 90),
    ("unknown.thing", {}, 50),
])
def test_every_rule_row(probe, evidence, points):
    assert finding_agility(probe, evidence) == points
