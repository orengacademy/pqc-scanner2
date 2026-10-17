from pqcscan.core.parity import drift_rows, parity_key
from pqcscan.core.types import Classification, Finding, Severity
from pqcscan.store.bundle import HostIdentity, dump_bundle, load_bundle
from pqcscan.store.repo import Repo


def test_parity_key_strips_endpoints():
    k1 = parity_key("net.tls.https", "RSA-2048", "server at 10.0.0.1:443 offers RSA-2048", {"endpoint": "10.0.0.1:443"})
    k2 = parity_key("net.tls.https", "RSA-2048", "server at 10.0.0.2:443 offers RSA-2048", {"endpoint": "10.0.0.2:443"})
    assert k1 == k2 and ":443" in k1[2]
    assert parity_key("ot.modbus.tcp", "plain-Modbus", "plain Modbus/TCP at plc-a:502", {}) == \
        parity_key("ot.modbus.tcp", "plain-Modbus", "plain Modbus/TCP at [fe80::1]:502", {})
    assert parity_key("fs.cert.x509", "RSA-2048", "cert /etc/ssl/a.pem", {"path": "/etc/ssl/a.pem"}) != \
        parity_key("fs.cert.x509", "RSA-2048", "cert /etc/ssl/b.pem", {"path": "/etc/ssl/b.pem"})


def test_parity_key_host_literals_and_server_attribution():
    # evidence["host"] is the scan target; bare IPv4/IPv6 literals and the host name drop out of titles
    a = parity_key("net.ssh", "ssh-rsa", "web01 offers ssh-rsa on 10.1.2.3", {"host": "web01", "_fleet": {"hosts": 1}})
    b = parity_key("net.ssh", "ssh-rsa", "web02 offers ssh-rsa on fe80::2", {"host": "web02"})
    assert a == b
    # same port on one endpoint stays distinct from another port
    assert parity_key("net.tls.https", "RSA-2048", "weak", {"endpoint": "h:443"}) != \
        parity_key("net.tls.https", "RSA-2048", "weak", {"endpoint": "h:8443"})


def _f(alg, title, probe="fs.cert.x509"):
    return Finding(probe_id=probe, algorithm=alg, classification=Classification.TINGGI,
                   severity=Severity.HIGH, title=title)


def test_drift_rows_added_and_removed(tmp_db_path):
    repo = Repo(tmp_db_path); repo.init_schema()
    a = repo.create_scan(mode="user", probe_versions={}, tool_versions={})
    b = repo.create_scan(mode="user", probe_versions={}, tool_versions={})
    repo.record_finding(a, _f("RSA-2048", "x")); repo.record_finding(a, _f("RSA-1024", "y"))
    repo.record_finding(b, _f("RSA-2048", "x")); repo.record_finding(b, _f("ECDSA-P256", "z"))
    rows = drift_rows(repo.list_findings(a), repo.list_findings(b))
    changes = sorted((r.evidence["change"], r.evidence["algorithm"]) for r in rows)
    assert changes == [("added", "ECDSA-P256"), ("removed", "RSA-1024")]
    assert all(r.probe_id == "fleet.drift" and r.classification is Classification.INFO for r in rows)


def test_drift_rows_ignore_prior_drift(tmp_db_path):
    repo = Repo(tmp_db_path); repo.init_schema()
    a = repo.create_scan(mode="user", probe_versions={}, tool_versions={})
    b = repo.create_scan(mode="user", probe_versions={}, tool_versions={})
    repo.record_finding(a, _f("RSA-2048", "x")); repo.record_finding(b, _f("RSA-2048", "x"))
    repo.record_finding(a, _f("RSA-1024", "old", probe="fleet.drift"))
    assert drift_rows(repo.list_findings(a), repo.list_findings(b)) == []


def test_ingest_emits_drift_against_previous_scan(tmp_db_path, tmp_path):
    src = Repo(tmp_path / "s.db"); src.init_schema()
    fleet = Repo(tmp_db_path); fleet.init_schema()
    ident = HostIdentity(machine_id="m", hostname="h", os="linux")
    h = fleet.upsert_host(ident)
    for algs in (["RSA-2048"], ["RSA-2048", "RSA-1024"]):
        sid = src.create_scan(mode="user", probe_versions={}, tool_versions={})
        for a in algs:
            src.record_finding(sid, _f(a, f"cert {a}"))
        src.finish_scan(sid, status="done")
        res = fleet.ingest_bundle(load_bundle(dump_bundle(src, sid, ident)), h.id)
    drift = [r for r in fleet.list_findings(res.scan_id) if r.probe_id == "fleet.drift"]
    assert len(drift) == 1 and drift[0].evidence["change"] == "added" and drift[0].evidence["algorithm"] == "RSA-1024"
    d = fleet.diff_findings(current_scan_id=res.scan_id, baseline_scan_id=res.scan_id - 1, parity=True)
    assert len(d["added"]) == 1          # drift rows themselves are excluded from parity diffs


def test_diff_cli_parity_ignores_host_names(tmp_db_path):
    from click.testing import CliRunner

    from pqcscan.cli.fleet import diff_cmd
    repo = Repo(tmp_db_path); repo.init_schema()
    ids = []
    for host in ("10.0.0.1", "10.0.0.2"):
        sid = repo.create_scan(mode="user", probe_versions={}, tool_versions={})
        repo.record_finding(sid, Finding(probe_id="net.tls.https", algorithm="RSA-2048",
                                         classification=Classification.TINGGI, severity=Severity.HIGH,
                                         title=f"server at {host}:443 offers RSA-2048",
                                         evidence={"host": host, "endpoint": f"{host}:443"}))
        ids.append(sid)
    args = ["--db", str(tmp_db_path), "--scan", str(ids[1]), "--baseline-scan", str(ids[0]), "--fail-on-drift"]
    assert CliRunner().invoke(diff_cmd, args).exit_code == 1
    assert CliRunner().invoke(diff_cmd, [*args, "--parity"]).exit_code == 0


def test_parity_key_is_linear_on_hostile_titles():
    import time
    for title in ("a:1" * 17_000, "a" * 50_000 + ":", "1.1.1.1" * 7_000, "[" + "f:" * 25_000):
        t0 = time.perf_counter()
        key = parity_key("net.tls.https", "RSA-2048", title, {"host": "a" * 300, "endpoint": title})
        assert time.perf_counter() - t0 < 0.5
        assert len(key[2]) < 700


def test_plain_diff_and_cli_ignore_drift_rows(tmp_db_path):
    from click.testing import CliRunner

    from pqcscan.cli.fleet import diff_cmd
    repo = Repo(tmp_db_path); repo.init_schema()
    s2 = repo.create_scan(mode="user", probe_versions={}, tool_versions={})
    s3 = repo.create_scan(mode="user", probe_versions={}, tool_versions={})
    for sid in (s2, s3):
        repo.record_finding(sid, _f("RSA-2048", "x"))
    repo.record_finding(s2, _f("RSA-1024", "gone", probe="fleet.drift"))
    for kw in ({}, {"parity": True}):
        d = repo.diff_findings(current_scan_id=s3, baseline_scan_id=s2, **kw)
        assert d["added"] == [] and d["removed"] == [] and d["common"] == 1
    args = ["--db", str(tmp_db_path), "--scan", str(s3), "--baseline-scan", str(s2), "--fail-on-drift"]
    assert CliRunner().invoke(diff_cmd, args).exit_code == 0


def test_ingest_never_stores_agent_sent_drift(tmp_db_path, tmp_path):
    src = Repo(tmp_path / "s.db"); src.init_schema()
    fleet = Repo(tmp_db_path); fleet.init_schema()
    ident = HostIdentity(machine_id="m", hostname="h", os="linux")
    h = fleet.upsert_host(ident)
    for extra in (False, True):
        sid = src.create_scan(mode="user", probe_versions={}, tool_versions={})
        src.record_finding(sid, _f("RSA-2048", "cert"))
        if extra:
            src.record_finding(sid, _f("RSA-1024", "forged", probe="fleet.drift"))
        src.finish_scan(sid, status="done")
        res = fleet.ingest_bundle(load_bundle(dump_bundle(src, sid, ident)), h.id)
    assert res.drift == 0 and res.stored_findings == 1
    assert [r.probe_id for r in fleet.list_findings(res.scan_id)] == ["fs.cert.x509"]


def test_parity_diff_cli_fails_on_real_difference(tmp_db_path):
    from click.testing import CliRunner

    from pqcscan.cli.fleet import diff_cmd
    repo = Repo(tmp_db_path); repo.init_schema()
    a = repo.create_scan(mode="user", probe_versions={}, tool_versions={})
    b = repo.create_scan(mode="user", probe_versions={}, tool_versions={})
    repo.record_finding(a, _f("RSA-2048", "x")); repo.record_finding(b, _f("RSA-1024", "x"))
    args = ["--db", str(tmp_db_path), "--scan", str(b), "--baseline-scan", str(a), "--parity", "--fail-on-drift"]
    assert CliRunner().invoke(diff_cmd, args).exit_code == 1


def test_empty_scan_after_findings_reports_every_removal(tmp_db_path):
    repo = Repo(tmp_db_path); repo.init_schema()
    a = repo.create_scan(mode="user", probe_versions={}, tool_versions={})
    b = repo.create_scan(mode="user", probe_versions={}, tool_versions={})
    for i in range(4):
        repo.record_finding(a, _f("RSA-2048", f"cert {i}"))
    rows = drift_rows(repo.list_findings(a), repo.list_findings(b))
    assert [r.evidence["change"] for r in rows] == ["removed"] * 4


def test_parity_key_caps_title_and_path_at_512():
    head = "x" * 512
    assert parity_key("p", "A", head + "tail-one", {}) == parity_key("p", "A", head + "tail-two", {})
    assert parity_key("p", "A", "t", {"path": "/" + "d" * 600 + "1"}) == \
        parity_key("p", "A", "t", {"path": "/" + "d" * 600 + "2"})
    long_title = "y" * 600
    d = drift_rows([], [type("R", (), {"probe_id": "p", "algorithm": "A", "title": long_title, "evidence": {}})()])
    assert len(d[0].evidence["title"]) == 512


def test_parity_key_normalises_at_host_and_bare_ipv6():
    assert parity_key("p", "A", "listener at web01.example.org", {}) == parity_key("p", "A", "listener at dr-02", {})
    assert parity_key("p", "A", "peer 2001:db8::1 uses RSA", {}) == parity_key("p", "A", "peer fe80::2 uses RSA", {})
    assert parity_key("p", "A", "peer (10.0.0.9), ok", {}) == parity_key("p", "A", "peer (10.0.0.8), ok", {})


def test_drift_rows_never_count_toward_readiness_or_totals(tmp_db_path, tmp_path, fast_registry):
    # scan 1 = RSA + ML-KEM, scan 2 = RSA only: the "removed ML-KEM" drift marker must not read as a green asset
    from fastapi.testclient import TestClient

    from pqcscan.core.bands import classify_band
    from pqcscan.daemon.app import create_app
    from pqcscan.renderers._report_context import build_report_context
    app = create_app(db_path=tmp_db_path, registry=fast_registry)
    fleet = app.state.repo
    src = Repo(tmp_path / "s.db"); src.init_schema()
    ident = HostIdentity(machine_id="m", hostname="h", os="linux")
    h = fleet.upsert_host(ident)
    for algs in (["RSA-2048", "ML-KEM-768"], ["RSA-2048"]):
        sid = src.create_scan(mode="user", probe_versions={}, tool_versions={})
        for a in algs:
            cls = Classification.PQC_READY if a.startswith("ML-") else Classification.TINGGI
            src.record_finding(sid, Finding(probe_id="net.tls.https", algorithm=a, classification=cls,
                                            severity=Severity.HIGH, title=f"server offers {a}"))
        src.finish_scan(sid, status="done")
        res = fleet.ingest_bundle(load_bundle(dump_bundle(src, sid, ident)), h.id)
    (d,) = [r for r in fleet.list_findings(res.scan_id) if r.probe_id == "fleet.drift"]
    assert d.algorithm == "N/A" and d.evidence["algorithm"] == "ML-KEM-768" and classify_band(d) == "grey"
    ctx = build_report_context(fleet, res.scan_id)
    assert (ctx["readiness"], ctx["bands"]["total"], ctx["total_findings"]) == (50, 1, 1)
    html = TestClient(app, base_url="http://127.0.0.1").get(f"/scans/{res.scan_id}").text
    assert "50/100" in html and "75/100" not in html
    assert ">1 " in html and ">2 " not in html
    from openpyxl import load_workbook

    from pqcscan.renderers.xlsx_bukukerja import render_xlsx_bukukerja
    render_xlsx_bukukerja(fleet, res.scan_id, tmp_path / "b.xlsx", locale="en")
    inv = [r for r in load_workbook(tmp_path / "b.xlsx")["0_Inventory"].iter_rows(min_row=5, values_only=True) if r[0]]
    assert len(inv) == 1 and inv[0][5] == "RSA-2048"       # no inventory asset made from the drift marker
